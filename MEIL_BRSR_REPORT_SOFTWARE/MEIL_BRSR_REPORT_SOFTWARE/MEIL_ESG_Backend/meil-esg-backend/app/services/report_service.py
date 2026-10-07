"""
Mirrors the frontend's lib/exporters.js exactly in spirit: every export is explicitly labelled a
*draft* (critique §36 — "Official BRSR report" overclaimed what a frontend-generated file is),
the XML is explicitly "XBRL-style, not filing-ready", and the report "package" (Comprehensive vs.
BRSR-Core-only, value-chain include/not-applicable, assessment/assurance/not-applicable) is a
request parameter, not a hard-coded assumption.
"""
import csv
import io
import json
from datetime import datetime, timezone

from app.models.brsr import BrsrPrinciple, BrsrCoreAttribute
from app.services import brsr_service


def build_report_model(*, scope_label, boundary_label, package, include_value_chain, disclosure_type,
                        fy, prev_fy, totals, prev_totals, kpi, prev_kpi, entity_disclosure,
                        core_snapshot, value_chain_summary, turnover_cr, projects, reporting_sites):
    principles = []
    for p in BrsrPrinciple.query.order_by(BrsrPrinciple.sequence).all():
        rows = []
        for ind in p.indicators:
            if package == "CORE" and not ind.core_attribute_id:
                continue
            rows.append({
                "code": ind.indicator_code, "label": ind.label, "unit": ind.unit,
                "tier": ind.tier, "core": ind.core_attribute_id,
                "current": totals.get(ind.indicator_code, 0), "previous": prev_totals.get(ind.indicator_code, 0),
            })
        if rows:
            principles.append({"code": p.code, "label": p.label, "rows": rows})

    return {
        "meta": {
            "scopeLabel": scope_label, "boundary": boundary_label, "fy": fy, "previousFy": prev_fy,
            "package": package, "includeValueChain": include_value_chain, "disclosureType": disclosure_type,
            "generatedAt": datetime.now(timezone.utc).isoformat(),
            "projects": projects, "reportingSites": reporting_sites, "turnoverCr": turnover_cr,
            "isDraft": True,
        },
        "sectionA": _section_a(entity_disclosure), "sectionB": _section_b(entity_disclosure),
        "principles": principles, "core": core_snapshot, "kpi": kpi, "kpiPrevious": prev_kpi,
        "valueChain": value_chain_summary if include_value_chain else None,
    }


def _section_a(ed):
    if ed is None:
        return {}
    return {
        "entityName": ed.entity_name, "cin": ed.cin, "registeredOffice": ed.registered_office,
        "boundaryDeclared": ed.boundary_declared, "csrApplicable": ed.csr_applicable,
        "stakeholderComplaints": ed.stakeholder_complaints_json or {},
    }


def _section_b(ed):
    if ed is None:
        return {}
    return {
        "directorStatement": ed.director_statement, "esgHead": ed.esg_head, "committee": ed.board_committee,
        "policies": [
            {"principle": pc.principle.code, "hasPolicy": pc.has_policy, "boardApproved": pc.board_approved,
             "extendsToValueChain": pc.extends_to_value_chain}
            for pc in ed.policy_coverage
        ] if ed.policy_coverage else [],
    }


def to_csv(model: dict) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([f"# {model['meta']['scopeLabel']} | {model['meta']['boundary']} | {model['meta']['fy']} | DRAFT — not a filed SEBI submission"])
    w.writerow(["Principle", "Indicator Code", "Indicator", "Tier", "Core Attribute", "Unit", model["meta"]["fy"], model["meta"]["previousFy"]])
    for p in model["principles"]:
        for r in p["rows"]:
            w.writerow([p["code"], r["code"], r["label"], r["tier"], r["core"] or "", r["unit"], r["current"], r["previous"]])
    w.writerow([])
    w.writerow(["BRSR Core attribute", "Metric", "Unit", "Current", "Status"])
    for c in model["core"]:
        for m in c["metrics"]:
            w.writerow([c["label"], m["label"], m.get("unit", ""), m.get("value"), c["disclosureStatus"]])
    return buf.getvalue()


def to_json(model: dict) -> str:
    return json.dumps(model, indent=2, default=str)


def to_xml(model: dict) -> str:
    def esc(s):
        return str(s if s is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")

    lines = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        "<!-- XBRL-STYLE DRAFT INSTANCE — NOT a filing-ready XBRL taxonomy instance. "
        "A filing-ready file needs the official SEBI/NSE BRSR taxonomy element mapping. -->",
        f'<BRSRDraft scope="{esc(model["meta"]["scopeLabel"])}" period="{esc(model["meta"]["fy"])}" '
        f'boundary="{esc(model["meta"]["boundary"])}" package="{esc(model["meta"]["package"])}">',
    ]
    for p in model["principles"]:
        lines.append(f'  <Principle code="{esc(p["code"])}" label="{esc(p["label"])}">')
        for r in p["rows"]:
            lines.append(f'    <Indicator code="{esc(r["code"])}" unit="{esc(r["unit"])}" tier="{esc(r["tier"])}">'
                          f'<Current>{r["current"]}</Current><Previous>{r["previous"]}</Previous></Indicator>')
        lines.append("  </Principle>")
    lines.append("  <BRSRCore>")
    for c in model["core"]:
        lines.append(f'    <Attribute id="{esc(c["id"])}" disclosureStatus="{esc(c["disclosureStatus"])}">')
        for m in c["metrics"]:
            lines.append(f'      <Metric label="{esc(m["label"])}" unit="{esc(m.get("unit",""))}">{m.get("value")}</Metric>')
        lines.append("    </Attribute>")
    lines.append("  </BRSRCore>")
    lines.append("</BRSRDraft>")
    return "\n".join(lines)


def to_pdf(model: dict) -> bytes:
    """Uses reportlab (already a hard dependency) rather than WeasyPrint, which needs system
    libraries (Pango/Cairo) that aren't guaranteed to be present in every deploy target — this
    keeps `pip install -r requirements.txt` sufficient with no OS packages required."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], fontSize=16)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], fontSize=12, spaceBefore=10)
    warn = ParagraphStyle("warn", parent=styles["Normal"], backColor=colors.HexColor("#FEF3C7"),
                           borderPadding=6, spaceAfter=10)
    small = styles["Normal"]

    story = [
        Paragraph(
            "<b>DRAFT / PREVIEW</b> — generated from approved site data. This is not a board-signed, "
            "filed BRSR report; the company secretary and statutory auditor should review it before filing.",
            warn,
        ),
        Paragraph("Business Responsibility &amp; Sustainability Report — Draft", h1),
        Paragraph(f"{model['meta']['scopeLabel']} · {model['meta']['fy']} · {model['meta']['boundary']} · "
                  f"Package: {model['meta']['package']}", small),
        Spacer(1, 10),
        Paragraph("BRSR Core", h2),
    ]
    core_rows = [["Attribute", "Metric", "Value", "Status"]]
    for c in model["core"]:
        for i, m in enumerate(c["metrics"]):
            core_rows.append([c["label"] if i == 0 else "", m["label"],
                               f'{m.get("value"):.2f}' if isinstance(m.get("value"), (int, float)) else "–",
                               c["disclosureStatus"] if i == 0 else ""])
    t = Table(core_rows, hAlign="LEFT", colWidths=[110, 220, 70, 90])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    story.append(t)

    for p in model["principles"]:
        story.append(Paragraph(f"{p['code']}: {p['label']}", h2))
        rows = [["Indicator", "Unit", model["meta"]["fy"], model["meta"]["previousFy"]]]
        for r in p["rows"]:
            rows.append([r["label"], r["unit"] or "", f'{r["current"]:.2f}' if isinstance(r["current"], (int, float)) else str(r["current"]),
                         f'{r["previous"]:.2f}' if isinstance(r["previous"], (int, float)) else str(r["previous"])])
        pt = Table(rows, hAlign="LEFT", colWidths=[260, 70, 80, 80])
        pt.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
        ]))
        story.append(pt)
        story.append(Spacer(1, 6))

    doc.build(story)
    return buf.getvalue()


def to_xlsx(model: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="E2E8F0")
    bold = Font(bold=True)

    ws = wb.active
    ws.title = "Summary"
    ws.append(["DRAFT / PREVIEW — not a filed BRSR report"])
    ws.append(["Scope", model["meta"]["scopeLabel"]])
    ws.append(["Boundary", model["meta"]["boundary"]])
    ws.append(["Financial Year", model["meta"]["fy"]])
    ws.append(["Package", model["meta"]["package"]])

    core_ws = wb.create_sheet("BRSR Core")
    core_ws.append(["Attribute", "Metric", "Unit", "Value", "Status"])
    for cell in core_ws[1]:
        cell.font = bold
        cell.fill = header_fill
    for c in model["core"]:
        for m in c["metrics"]:
            core_ws.append([c["label"], m["label"], m.get("unit", ""), m.get("value"), c["disclosureStatus"]])

    for p in model["principles"]:
        ws_p = wb.create_sheet(p["code"][:31])
        ws_p.append(["Indicator", "Unit", "Tier", model["meta"]["fy"], model["meta"]["previousFy"]])
        for cell in ws_p[1]:
            cell.font = bold
            cell.fill = header_fill
        for r in p["rows"]:
            ws_p.append([r["label"], r["unit"], r["tier"], r["current"], r["previous"]])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
